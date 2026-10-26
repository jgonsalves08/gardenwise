"""Exports (CODE, no AI cost): calendar file, printable plan, and a grant-ready budget sheet."""

from __future__ import annotations

import base64
import hashlib
import html
import io
import re
from datetime import date, datetime, timedelta, timezone

from . import config
from .data import plants
from .pipeline import Plan, task_key

PRODID = f"-//{config.APP_NAME}//Garden plan//EN"


# ---------- 1. calendar (.ics) ----------
def _ics_escape(text: str) -> str:
    return (text or "").replace("\\", "\\\\").replace(";", r"\;").replace(",", r"\,").replace("\n", r"\n")


def _fold(line: str) -> str:
    """RFC 5545: lines longer than 75 octets are folded with CRLF + space."""
    out, raw = [], line.encode("utf-8")
    while len(raw) > 75:
        cut = 75
        while cut > 0 and (raw[cut] & 0xC0) == 0x80:   # don't split a UTF-8 character
            cut -= 1
        out.append(raw[:cut].decode("utf-8"))
        raw = b" " + raw[cut:]
    out.append(raw.decode("utf-8"))
    return "\r\n".join(out)


def to_ics(plan: Plan, garden_name: str, done: dict | None = None, reminders: bool = True) -> bytes:
    """All-day events for every task (moved dates included), with a reminder at 9am the day before."""
    done = done or {}
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    lines = ["BEGIN:VCALENDAR", "VERSION:2.0", f"PRODID:{PRODID}", "CALSCALE:GREGORIAN", "METHOD:PUBLISH",
             f"X-WR-CALNAME:{_ics_escape(garden_name)} ({config.APP_NAME})"]
    for t in plan.tasks:
        uid = hashlib.sha1(f"{garden_name}|{task_key(t)}".encode()).hexdigest()[:20]
        status = " ✓" if task_key(t) in done else ""
        desc = t.detail + (f"\n(Gardenwise recommended {t.recommended:%b %d}.)" if t.recommended else "")
        lines += ["BEGIN:VEVENT", f"UID:{uid}@gardenwise", f"DTSTAMP:{stamp}",
                  f"DTSTART;VALUE=DATE:{t.when:%Y%m%d}", f"DTEND;VALUE=DATE:{t.when + timedelta(days=1):%Y%m%d}",
                  f"SUMMARY:{_ics_escape('🌱 ' + t.title + status)}",
                  f"DESCRIPTION:{_ics_escape(desc + f' — from your {config.APP_NAME} plan for {garden_name}.')}",
                  "TRANSP:TRANSPARENT"]
        if reminders and task_key(t) not in done:
            lines += ["BEGIN:VALARM", "ACTION:DISPLAY", f"DESCRIPTION:{_ics_escape(t.title)}",
                      "TRIGGER:-PT15H", "END:VALARM"]   # 9am the day before an all-day event
        lines.append("END:VEVENT")
    lines.append("END:VCALENDAR")
    return ("\r\n".join(_fold(x) for x in lines) + "\r\n").encode("utf-8")


# ---------- 2. printable plan (self-contained HTML; Print → Save as PDF) ----------
def _photo(pid: str) -> str:
    p = config.DATA_DIR / "plant_images" / f"{pid}.jpg"
    if not p.exists():
        return ""
    return f'<img src="data:image/jpeg;base64,{base64.b64encode(p.read_bytes()).decode()}" alt="">'


def _alert_html(text: str) -> str:
    """An alert's simple markdown (**bold**, "- " bullet lines, blank-line paragraphs) as safe HTML."""
    out, items = [], []
    for line in html.escape(text).split("\n"):
        line = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", line.strip())
        if line.startswith("- "):
            items.append(f"<li>{line[2:]}</li>")
            continue
        if items:
            out.append("<ul>" + "".join(items) + "</ul>")
            items = []
        if line:
            out.append(f"<div>{line}</div>")
    if items:
        out.append("<ul>" + "".join(items) + "</ul>")
    return "".join(out)


def printable_html(plan: Plan, garden_name: str, season: str = "", done: dict | None = None,
                   icons: dict | None = None) -> bytes:
    done, icons = done or {}, icons or {}
    c, loc, db = plan.constraints, plan.location, plants()
    alloc = {a.plant_id: a for a in plan.allocations}
    sched = {s.plant_id: s for s in plan.schedules}
    e = html.escape
    rows = ""
    for p in plan.plants:
        a, s = alloc.get(p.id), sched.get(p.id)
        rows += (f"<tr><td class='ph'>{_photo(p.id)}</td><td><b>{e(p.name)}</b><br><span class='m'>{e(p.tip)}</span></td>"
                 f"<td>{a.count if a else '—'}</td><td>{s.plant_out:%b %d, %Y}</td>"
                 f"<td>{s.harvest_start:%b %d} – {s.harvest_end:%b %d}</td></tr>" if s else "")
    grid = ""
    if plan.grid:
        grid = "<table class='grid'>" + "".join(
            "<tr>" + "".join(f"<td>{icons.get(cid, '🌱')}</td>" if cid else "<td class='free'></td>" if cid is None
                             else "<td class='off'></td>" for cid in row) + "</tr>"
            for row in plan.grid) + "</table>"
        grid += "<p class='m'>" + " · ".join(
            [f"{icons.get(a.plant_id, '🌱')} {e(db[a.plant_id].name)} × {a.count}" for a in plan.allocations if a.count]
            + ([f"➕ free space: {plan.free_space} sq ft (dashed squares), room for more plants"]
               if plan.free_space else [])) + "</p>"
    tasks, month = "", None
    for t in plan.tasks:
        m = t.when.strftime("%B %Y")
        if m != month:
            tasks += f"<h3>{m}</h3>"
            month = m
        box = "☑" if task_key(t) in done else "☐"
        tasks += f"<div class='t'>{box} <b>{t.when:%b %d}</b> · {e(t.title)}</div>"
    shop = ""
    for group, m in plan.shopping_lines():
        shop += f"<tr><td>☐ {e(m.item)}</td><td>{m.qty}</td><td>~${round(m.subtotal):,}</td></tr>"
    warn = "".join(f"<li>{_alert_html(w)}</li>" for w in plan.warnings)
    doc = f"""<!doctype html><html><head><meta charset="utf-8"><title>{e(garden_name)} · {config.APP_NAME}</title>
<style>
html {{color-scheme: light; background: #FFFFFF;}}
body {{font-family: Georgia, 'Times New Roman', serif; color:#1E2A22; background: #FFFFFF; max-width: 820px; margin: 24px auto; padding: 0 16px;}}
h1 {{color:#2F6B3F; margin-bottom: 0;}} h2 {{color:#2F6B3F; border-bottom: 2px solid #D6E6D3; padding-bottom: 4px; margin-top: 28px;}}
h3 {{margin: 14px 0 4px 0; color:#4E7A55;}} .m {{color:#5C6B5F; font-size: 13px;}}
table {{border-collapse: collapse; width: 100%;}} td, th {{padding: 6px 8px; border-bottom: 1px solid #E3EDE0; vertical-align: top; font-size: 14px;}}
th {{text-align: left; color:#2F6B3F;}} td.ph img {{width: 64px; height: 64px; object-fit: cover; border-radius: 8px;}}
.grid td {{width: 34px; height: 34px; text-align: center; border: 1px solid #fff; background: #EEF6EC; font-size: 18px;}}
.grid td.free {{background: #fff; border: 2px dashed #B9D3B2;}} .grid td.off {{background: none;}}
.grid {{width: auto;}} .t {{font-size: 14px; padding: 2px 0;}} .btn {{background:#2F6B3F; color:#fff; border:none; padding:8px 14px; border-radius:8px; cursor:pointer;}}
@media print {{ .noprint {{display:none;}} h2 {{page-break-after: avoid;}} }}
</style></head><body>
<p class="noprint"><button class="btn" onclick="window.print()">🖨️ Print or save as PDF</button></p>
<h1>🌱 {e(garden_name)}</h1>
<p class="m">{e(loc.label)} · {e(season)} · start {c.start_date:%b %d, %Y} · average frosts: last {loc.last_frost_on(2000):%b %d}, first {loc.first_frost_on(2000):%b %d}</p>
<h2>Plants</h2><table><tr><th></th><th>Plant</th><th>How many</th><th>Plant out</th><th>Harvest</th></tr>{rows}</table>
{('<h2>Bed layout (north at top, 1 square = 1 ft)</h2>' + grid) if grid else ''}
<h2>Calendar</h2>{tasks}
<h2>Shopping list (estimated total ~${round(plan.total):,})</h2>
<table><tr><th>Item</th><th>Qty</th><th>Est.</th></tr>{shop}</table>
<p class="m">{e(config.SHOP_PRICE_NOTE)}</p>
{('<h2>Things to know</h2><ul>' + warn + '</ul>') if warn else ''}
<p class="m">Made with {config.APP_NAME} on {date.today():%B %d, %Y}. Plant photos: Wikimedia Commons (see the app's About &amp; credits page).</p>
</body></html>"""
    return doc.encode("utf-8")


# ---------- 3. grant-ready budget sheet (Excel) ----------
def _safe(v):
    """Neutralize spreadsheet formula injection in any text cell."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


def budget_xlsx(title: str, sections: list[tuple[str, list[dict]]], summary: list[tuple[str, object]],
                notes: list[str]) -> bytes:
    """sections: [(sheet name, rows as dicts)]. First sheet is a summary for the grant application."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    green = PatternFill("solid", fgColor="2F6B3F")
    ws.append([_safe(title)])
    ws["A1"].font = Font(bold=True, size=14, color="2F6B3F")
    ws.append([f"Prepared with {config.APP_NAME} on {date.today():%B %d, %Y}"])
    ws.append([])
    for k, v in summary:
        ws.append([_safe(k), _safe(v)])
    ws.append([])
    for n in notes:
        ws.append([_safe(n)])
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 22
    for name, rows in sections:
        sh = wb.create_sheet(name[:31])
        if not rows:
            sh.append(["(none)"])
            continue
        cols = list(rows[0].keys())
        sh.append(cols)
        for cell in sh[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = green
            cell.alignment = Alignment(vertical="center")
        for r in rows:
            sh.append([_safe(r.get(c)) for c in cols])
        for i, c in enumerate(cols, start=1):
            sh.column_dimensions[sh.cell(1, i).column_letter].width = max(12, min(48, len(str(c)) + 6))
            if any(w in c.lower() for w in ("cost", "price", "each", "total", "$")):
                for row in sh.iter_rows(min_row=2, min_col=i, max_col=i):
                    for cell in row:
                        if isinstance(cell.value, (int, float)):
                            cell.number_format = '"$"#,##0.00'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def garden_budget(plan: Plan, garden_name: str, season: str = "") -> bytes:
    items = [{"Category": "Plants & seeds" if g == "plants" else next(b.name for b in plan.builds if b.build_id == g),
              "Item": m.item, "Qty": m.qty, "Each ($)": m.unit_price, "Cost ($)": m.subtotal}
             for g, m in plan.shopping_lines()]
    by_cat: dict[str, float] = {}
    for it in items:
        by_cat[it["Category"]] = by_cat.get(it["Category"], 0) + it["Cost ($)"]
    lb, value = plan.expected_harvest()
    c = plan.constraints
    summary = [("Garden", f"{garden_name} {season}".strip()), ("Location", plan.location.label),
               ("Garden for", {"me": "Individual", "household": f"Household ({c.adults} adults, {c.kids} kids)",
                               "school": f"School (grades {c.grade_band})", "community": "Community garden plot"}[c.audience]),
               ("Estimated total cost ($)", round(plan.total, 2))]
    summary += [(f"  {k} ($)", round(v, 2)) for k, v in by_cat.items()]
    summary += [("Expected harvest (approx. lb)", lb), ("Approx. grocery value of harvest ($)", round(value, 2)),
                ("Plants", ", ".join(p.name for p in plan.plants))]
    notes = [config.SHOP_PRICE_NOTE, config.GROCERY_PRICE_NOTE]
    return budget_xlsx(f"Garden budget: {garden_name}", [("Line items", items)], summary, notes)


def community_budget(group: dict, rollup) -> bytes:
    pooled = [{"Item": d["item"], "Plots using it": d["plots"], "Qty to buy": d["pooled_qty"],
               "Each ($)": d["unit_price"], "Cost ($)": d["pooled_cost"]} for d in rollup.combined]
    shared = [{"Build": b.name, "Item": m.item, "Qty": m.qty, "Each ($)": m.unit_price, "Cost ($)": m.subtotal}
              for b in rollup.shared_builds for m in b.materials]
    per_plot = [{"Plot": p["name"], "Own beds/pots ($)": p["own_builds"], "Plants & seeds ($)": p["plants"],
                 "Share of shared builds ($)": p["shared_share"], "Total ($)": p["total"]} for p in rollup.per_plot]
    summary = [("Community garden", group["name"]), ("Plots", len(rollup.per_plot)),
               ("Total budget, pooled purchasing ($)", rollup.pooled_total),
               ("If each plot bought separately ($)", rollup.separate_total),
               ("Savings from pooling ($)", round(rollup.separate_total - rollup.pooled_total, 2)),
               ("Average cost per plot ($)", round(rollup.pooled_total / max(1, len(rollup.per_plot)), 2))]
    notes = [config.SHOP_PRICE_NOTE,
             "Pooling assumes one seed packet covers about 3 plots and shared builds are bought once."]
    return budget_xlsx(f"Community garden budget: {group['name']}",
                       [("Pooled seed & plant order", pooled), ("Shared builds", shared), ("Cost per plot", per_plot)],
                       summary, notes)
